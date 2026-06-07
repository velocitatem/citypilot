"""Runs INSIDE the Daytona sandbox. The host invokes:

    python -m sandbox_lib <command>

with args at /tmp/args.json. Result is one JSON object on stdout.
"""

import json
import sys
from pathlib import Path


ROOT = Path("/home/daytona/workspace")
DB_PATH = ROOT / "data" / "agent.duckdb"
ARTIFACTS = ROOT / "artifacts"


def _ddb(read_only: bool = True):
    import duckdb
    return duckdb.connect(str(DB_PATH), read_only=read_only)


def list_tables():
    with _ddb() as c:
        return [r[0] for r in c.execute("SHOW TABLES").fetchall()]


def describe_table(table: str):
    with _ddb() as c:
        return c.execute(f"DESCRIBE {table}").fetchdf().to_dict("records")


def query_db(sql: str):
    s = sql.strip().rstrip(";")
    if ";" in s or not s.lstrip("(").lower().startswith(("select", "with")):
        raise ValueError("only single SELECT/WITH allowed")
    with _ddb() as c:
        return c.execute(s).fetchdf().to_dict("records")


def _artifact_path(kind: str, filename: str) -> Path:
    out = ARTIFACTS / kind / filename
    out.parent.mkdir(parents=True, exist_ok=True)
    return out


def create_pdf_report(filename: str, html: str):
    from weasyprint import HTML
    out = _artifact_path("pdf", filename)
    HTML(string=html).write_pdf(str(out))
    return str(out)


def create_docx_report(filename: str, sections: list[dict]):
    from docx import Document
    doc = Document()
    for sec in sections:
        if h := sec.get("heading"):
            doc.add_heading(h, level=sec.get("level", 1))
        if b := sec.get("body"):
            doc.add_paragraph(b)
    out = _artifact_path("docx", filename)
    doc.save(str(out))
    return str(out)


def create_xlsx_workbook(filename: str, sheets: dict[str, list[dict]]):
    import pandas as pd
    out = _artifact_path("xlsx", filename)
    with pd.ExcelWriter(out, engine="xlsxwriter") as w:
        for name, rows in sheets.items():
            pd.DataFrame(rows).to_excel(w, sheet_name=name[:31], index=False)
    return str(out)


def create_plotly_chart(filename: str, figure: dict):
    import plotly.graph_objects as go
    out = _artifact_path("plots", filename if filename.endswith(".html") else f"{filename}.html")
    fig = go.Figure(figure)
    fig.write_html(str(out), include_plotlyjs="cdn", full_html=True)
    return str(out)


def load_data_as_table(table_name: str, rows: list[dict]):
    # Sanitise name: only word characters allowed to prevent SQL injection.
    import re
    if not re.fullmatch(r"\w+", table_name):
        raise ValueError(f"Invalid table name: {table_name!r}")
    import pandas as pd
    df = pd.DataFrame(rows)
    with _ddb(read_only=False) as c:
        c.register("_incoming", df)
        c.execute(f'DROP TABLE IF EXISTS "{table_name}"')
        c.execute(f'CREATE TABLE "{table_name}" AS SELECT * FROM _incoming')
        count = c.execute(f'SELECT COUNT(*) FROM "{table_name}"').fetchone()[0]
    return {"table": table_name, "rows": count, "columns": list(df.columns)}


COMMANDS = {
    "list_tables": list_tables,
    "describe_table": describe_table,
    "query_db": query_db,
    "create_pdf_report": create_pdf_report,
    "create_docx_report": create_docx_report,
    "create_xlsx_workbook": create_xlsx_workbook,
    "create_plotly_chart": create_plotly_chart,
    "load_data_as_table": load_data_as_table,
}


if __name__ == "__main__":
    name = sys.argv[1]
    args = json.loads(Path("/tmp/args.json").read_text()) if Path("/tmp/args.json").exists() else {}
    try:
        result = COMMANDS[name](**args)
        sys.stdout.write(json.dumps({"ok": True, "result": result}, default=str))
    except Exception as e:
        sys.stdout.write(json.dumps({"ok": False, "error": f"{type(e).__name__}: {e}"}))
        sys.exit(1)
